"""Generate a deliberately messy supplier quotation workbook for the import demo.

Usage (from the repository root, database already seeded):
    python scripts/make_supplier_excel.py

seed_demo calls write_supplier_excel() directly. Layout follows
docs/data-dictionary.html §12: two company header rows before the real header,
80 data rows of which 10 repeat catalog OE numbers, 5 carry dirty numbers,
3 have no OE number and 2 have a currency symbol in the price.
"""

import datetime as dt
import random
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

HEADERS = [
    "Part No.", "OEM NO", "Ref", "Description", "适用车型", "MOQ", "FOB Price(USD)", "Packing",
]
DATA_ROWS = 80
FIXED_TIME = dt.datetime(2026, 9, 1, 9, 0, 0)

# Messy variants of what a real supplier sheet contains (§10).
VEHICLE_TEXT = {
    "Volvo": ["Volvo FH12/FH16", "VOLVO FM", "沃尔沃 FH"],
    "Scania": ["SCANIA R", "Scania P/G/R"],
    "Mercedes-Benz": ["Benz Actros", "奔驰 Axor", "MB Actros MP2"],
    "MAN": ["MAN TGA", "M.A.N TGX"],
    "DAF": ["DAF XF105", "daf CF85"],
    "Sinotruk HOWO": ["HOWO A7", "中国重汽 豪沃"],
}
NAMES = [
    ("Brake Pad Set", "刹车片"), ("Brake Disc", "刹车盘"), ("Air Filter", "空气滤清器"),
    ("Oil Filter", "机油滤清器"), ("Fuel Filter", "燃油滤清器"), ("Clutch Disc", "离合器片"),
    ("Shock Absorber", "减震器"), ("Water Pump", "水泵"),
]
CROSS_BRANDS = ["Knorr-Bremse", "WABCO", "MANN-FILTER", "Fleetguard", "Donaldson", "Sachs"]


def _moq(rng):
    n = rng.choice([10, 20, 50, 100, 200])
    return rng.choice([n, f"{n}pcs", f"{n} sets", str(n)])


def _packing(rng):
    return rng.choice([
        f"{rng.choice([1, 2, 4, 10, 20])}pcs/ctn",
        "1 set/box",
        f"{rng.randint(20, 60)}*{rng.randint(15, 40)}*{rng.randint(10, 35)}cm",
        f"{rng.randint(1, 20)}.{rng.randint(0, 9)}KGS",
    ])


def build_rows(samples: list[dict], taken: set[str], seed: int) -> list[list]:
    """Return 80 data rows. samples: catalog parts to repeat ({oe, name_en, vehicle})."""
    from apps.catalog.demo import NumberFactory, to_fullwidth

    rng = random.Random(seed + 1)
    numbers = NumberFactory(rng, taken=set(taken))
    makes = list(VEHICLE_TEXT)

    def base_row(oe, name, vehicle):
        return [
            f"HB-{rng.randint(1000, 9999)}",
            oe,
            numbers.new(rng.choice(CROSS_BRANDS)) if rng.random() < 0.5 else "",
            name,
            vehicle,
            _moq(rng),
            round(rng.uniform(3, 180), 2),
            _packing(rng),
        ]

    def new_row():
        make = rng.choice(makes)
        name_en, name_zh = rng.choice(NAMES)
        name = name_en if rng.random() > 0.1 else f"{name_zh} {name_en}"  # mixed zh/en
        return base_row(numbers.new(make), name, rng.choice(VEHICLE_TEXT[make]))

    rows = []
    for sample in samples[:10]:  # 10 rows repeat catalog parts
        rows.append(base_row(sample["oe"], sample["name_en"], sample["vehicle"]))

    dirty = [
        lambda: to_fullwidth(numbers.new("Volvo")),
        lambda: numbers.new("Sinotruk HOWO").lower(),
        lambda: f"{numbers.new('Volvo')} / {numbers.new('Volvo')}",
        lambda: f"{numbers.new('Scania')}\n{numbers.new('Scania')}",
        lambda: "  " + " ".join(numbers.new("DAF")) + " ",
    ]
    for make_dirty in dirty:  # 5 rows with dirty numbers
        row = new_row()
        row[1] = make_dirty()
        rows.append(row)

    for _ in range(3):  # 3 rows without an OE number
        row = new_row()
        row[1] = ""
        rows.append(row)

    while len(rows) < DATA_ROWS:
        rows.append(new_row())

    for i, label in [(-1, "USD {:.2f}"), (-2, "${:.2f}")]:  # 2 prices with currency symbol
        rows[i][6] = label.format(rows[i][6])

    rng.shuffle(rows)
    return rows


def write_supplier_excel(path: Path, samples: list[dict], taken: set[str], seed: int) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Quotation"
    ws.append(["Hebei Demo Auto Parts Co., Ltd.（演示用虚构供应商）"])
    ws.append([f"QUOTATION   Date: {FIXED_TIME:%Y-%m-%d}   Valid: 30 days"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(HEADERS))
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(HEADERS)
    for cell in ws[3]:
        cell.font = Font(bold=True)
    for row in build_rows(samples, taken, seed):
        ws.append(row)
    for col, width in zip("ABCDEFGH", [10, 22, 16, 30, 20, 10, 16, 16], strict=True):
        ws.column_dimensions[col].width = width

    wb.properties.created = wb.properties.modified = FIXED_TIME
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def main() -> None:
    import os

    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
    os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings.local")
    import django

    django.setup()
    from django.conf import settings

    from apps.catalog.demo import DEFAULT_SEED
    from apps.catalog.models import PartNumber

    oe = list(
        PartNumber.objects.filter(kind="OE").select_related("part").order_by("id")[:400]
    )
    if not oe:
        sys.exit("No OE numbers in the database; run `python manage.py seed_demo` first.")
    rng = random.Random(DEFAULT_SEED)
    samples = [
        {"oe": pn.number, "name_en": pn.part.name_en, "vehicle": ""}
        for pn in rng.sample(oe, k=min(10, len(oe)))
    ]
    taken = set(PartNumber.objects.values_list("number_norm", flat=True))
    path = Path(settings.MEDIA_ROOT) / "demo" / "supplier_quote_messy.xlsx"
    write_supplier_excel(path, samples, taken, DEFAULT_SEED)
    print(f"wrote {path}")


if __name__ == "__main__":
    main()
