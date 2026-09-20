import json
from io import StringIO
from pathlib import Path

import pytest
from django.core.management import CommandError, call_command
from django.db.models import Count
from openpyxl import load_workbook

from apps.catalog.models import Brand, Category, Fitment, Part, PartImage, PartNumber
from apps.catalog.services.normalize import normalize_number
from apps.suppliers.models import Supplier, SupplierOffer

pytestmark = pytest.mark.django_db


def seed(*args):
    call_command("seed_demo", *args, stdout=StringIO())


def snapshot():
    """Everything that must be identical between two runs with the same seed."""
    return [
        (
            p.sku, p.name_en, p.category.name if p.category else None, p.status, p.attributes,
            p.packaging, p.description_en,
            sorted((n.number, n.kind, n.brand.name if n.brand else None) for n in p.numbers.all()),
            sorted((f.make, f.model, f.engine, f.year_from, f.year_to) for f in p.fitments.all()),
            sorted((i.image.name, i.is_primary) for i in p.images.all()),
            sorted(
                (o.supplier.name, o.supplier_pn, o.unit_cost_usd, o.moq, o.lead_days, o.quoted_at)
                for o in p.offers.all()
            ),
        )
        for p in Part.objects.select_related("category").prefetch_related(
            "numbers__brand", "fitments", "images", "offers__supplier"
        )
    ]


def test_seed_demo_volume_and_dirty_data(settings):
    seed()

    parts = Part.objects.count()
    assert parts == 400
    assert Category.objects.filter(parent__isnull=False).count() == 12
    assert Brand.objects.filter(kind="oem").count() == 10
    assert Brand.objects.filter(kind="aftermarket").count() == 8
    # ~900 OE / cross numbers + one supplier number per quote (~500, added in M09)
    assert 1200 <= PartNumber.objects.count() <= 1650
    assert 450 <= SupplierOffer.objects.count() <= 650
    assert Supplier.objects.count() == 6
    assert 450 <= Fitment.objects.count() <= 700

    def share(item):
        return Part.objects.missing(item).count() / parts

    assert 0.28 <= share("image") <= 0.42
    assert 0.18 <= share("oe") <= 0.32
    assert 0.14 <= share("category") <= 0.26
    assert 0.23 <= share("fitment") <= 0.37
    assert 0.09 <= share("description") <= 0.21
    assert 0.18 <= share("offer") <= 0.32

    numbers = list(PartNumber.objects.values_list("number", flat=True))
    assert any("-" in n for n in numbers)
    assert any(any(ord(c) > 0xFF00 for c in n) for n in numbers)  # full-width
    assert any(any(c.islower() for c in n) for n in numbers)

    shared_oe = (
        PartNumber.objects.filter(kind="OE")
        .values("number_norm")
        .annotate(parts=Count("part", distinct=True))
        .filter(parts__gt=1)
    )
    assert shared_oe.count() >= 10

    statuses = set(Part.objects.values_list("status", flat=True))
    assert statuses == {"draft", "reviewed", "published"}


def test_every_generated_part_is_valid():
    seed("--parts", "80")

    for part in Part.objects.select_related("category"):
        part.full_clean()


def test_same_seed_gives_same_data():
    seed("--parts", "60")
    first = snapshot()

    seed("--parts", "60", "--reset")

    assert snapshot() == first


def test_refuses_to_overwrite_without_reset():
    seed("--parts", "5")

    with pytest.raises(CommandError):
        seed("--parts", "5")


def test_reset_leaves_no_orphan_image_files(settings):
    seed("--parts", "40")
    seed("--parts", "40", "--reset")

    files = [p for p in (Path(settings.MEDIA_ROOT) / "parts").rglob("*.jpg")]
    assert len(files) == PartImage.objects.count()


def test_supplier_excel_layout(settings):
    seed()

    path = Path(settings.MEDIA_ROOT) / "demo" / "supplier_quote_messy.xlsx"
    rows = list(load_workbook(path).active.iter_rows(values_only=True))
    header, data = rows[2], rows[3:]

    assert rows[0][0] and rows[1][0]  # two company rows above the header
    assert header == (
        "Part No.", "OEM NO", "Ref", "Description", "适用车型", "MOQ", "FOB Price(USD)", "Packing",
    )
    assert len(data) == 80

    catalog = set(PartNumber.objects.values_list("number_norm", flat=True))
    assert sum(1 for r in data if r[1] and normalize_number(r[1]) in catalog) == 10
    assert sum(1 for r in data if not r[1]) == 3
    assert sum(1 for r in data if isinstance(r[6], str)) == 2


def test_the_photo_demo_lands_on_a_part_that_can_be_quoted(settings, tmp_path):
    """DEMO_SCRIPT step 4 goes photo -> confirm -> quote, so the part the offline
    vision example points at must have a supplier offer."""
    settings.MEDIA_ROOT = tmp_path
    from apps.ai.prompts.loader import example_output
    from apps.suppliers.demo import SHOWCASE_OE

    example = json.loads(example_output("image_identify"))
    assert normalize_number(example["visible_numbers"][0]) == SHOWCASE_OE

    seed()

    part = Part.objects.get(numbers__number_norm=SHOWCASE_OE, numbers__kind="OE")
    assert part.offers.exists()


def test_reset_after_a_demo_run_with_quotations(settings, tmp_path, django_user_model):
    """A demo leaves inquiries and quote lines behind; QuoteLine PROTECTs its part,
    so --reset has to clear them first (DEMO_SCRIPT's "回到初始状态" fallback)."""
    from decimal import Decimal

    from apps.inquiries.models import Inquiry, QuoteLine
    from apps.suppliers.services.quoting import best_offer

    settings.MEDIA_ROOT = tmp_path
    seed()
    part = Part.objects.filter(offers__isnull=False).first()
    inquiry = Inquiry.objects.create(input_type="text", raw_input="4020269", matched_part=part)
    QuoteLine.objects.create(inquiry=inquiry, part=part, offer=best_offer(part), qty=50,
                             unit_cost_usd=Decimal("22.90"), margin=Decimal("0.250"),
                             unit_price_usd=Decimal("28.63"))

    seed("--reset")

    assert Part.objects.count() == 400
    assert not Inquiry.objects.exists() and not QuoteLine.objects.exists()
