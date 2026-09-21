"""Synthetic supplier files for tests. Two header styles that share no column name,
like two real suppliers would; every value is invented."""

import datetime as dt
import io

from openpyxl import Workbook

HEADERS_A = ["Source Record ID", "Supplier SKU", "Product Description", "Truck Application",
             "Position", "OE / Reference", "Package Size", "Unit Price", "Currency", "MOQ",
             "Quote Date"]
HEADERS_B = ["Source Ref", "Brand Number", "English Name", "Fitment", "Install Side",
             "Cross Ref", "Package", "Offer Price", "Curr.", "Min Order", "Quoted On"]
ROWS = [
    ["X-001", "X-01-00", "Front Grille", "Volvo VNL 2018-2023", None, "OE-TST-1001",
     "122 x 75 x 7 cm", 42.67, "USD", 2, dt.datetime(2026, 8, 3)],
    ["X-002", "X-02-L", "Side Grille, Left side", "Volvo VNL 2018-2023", "Left", "OE-TST-1002",
     "86 x 36 x 56 cm", 51.01, "EUR", 10, dt.datetime(2026, 8, 5)],
    ["X-003", "X-03-00", "Fan Shroud", "Freightliner Century", None, None,
     "121 x 22 x 92 cm", 80.2, "USD", 5, dt.datetime(2026, 8, 12)],
]


def workbook_bytes(headers=HEADERS_A, rows=ROWS, *, title_rows=(), sheet="Quote",
                   extra_sheets=(), footer=None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for text in title_rows:  # company banner lines above the real header
        ws.append([text])
    ws.append(headers)
    for row in rows:
        ws.append(row)
        if len(headers) > 7:  # the price column is formatted "$" whatever the currency
            ws.cell(ws.max_row, 8).number_format = "$#,##0.00"
    if footer:
        ws.append([])
        ws.append([footer])
    for name in extra_sheets:
        wb.create_sheet(name)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
