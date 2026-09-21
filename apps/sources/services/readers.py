"""Read supplier files cell by cell, keeping where every value came from
(docs/archive-design.html §5 step 2, §6).

Excel is read with openpyxl rather than pandas: lineage needs each cell's
coordinate, and the price check in M24 needs its number format.
"""

import csv
import datetime as dt
import io
from dataclasses import dataclass, field
from decimal import Decimal

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

HEADER_SCAN_ROWS = 20  # the header is looked for among the first rows that have content
MIN_CELLS_PER_ROW = 2  # rows with fewer filled cells are titles, notes or totals
CSV_ENCODINGS = ("utf-8-sig", "gb18030")


class ReadError(ValueError):
    """The file cannot be read; the message is shown to the user and stored on the file."""


@dataclass(frozen=True)
class Cell:
    value: object  # int, float, datetime, str ... as stored in the file
    text: str  # the stored value as text; never the display format
    coord: str  # "G14"
    number_format: str = ""


@dataclass
class Row:
    row_no: int
    cells: dict[str, Cell]  # header -> cell

    def texts(self) -> dict[str, str]:
        return {header: cell.text for header, cell in self.cells.items()}


@dataclass
class Sheet:
    name: str
    header_row: int
    headers: list[str]
    rows: list[Row]
    skipped: list[tuple[int, str]] = field(default_factory=list)  # (row_no, why)


def read_source(source_file) -> list[Sheet]:
    readers = {"xlsx": read_workbook, "csv": read_csv}
    reader = readers.get(source_file.file_format)
    if reader is None:
        raise ReadError(f"暂不支持读取 {source_file.get_file_format_display()} 文件。")
    with source_file.file.open("rb") as fh:
        sheets = reader(fh)
    if not sheets:
        raise ReadError("文件里没有找到带表头的数据表。")
    return sheets


def read_workbook(fh) -> list[Sheet]:
    try:
        workbook = load_workbook(fh, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many types for damaged files
        raise ReadError(f"无法读取 Excel 文件：{exc}") from exc
    sheets = []
    for worksheet in workbook.worksheets:
        grid = [
            [_cell(c, row_no, col_no) for col_no, c in enumerate(row, start=1)]
            for row_no, row in enumerate(worksheet.iter_rows(), start=1)
        ]
        sheet = _build_sheet(worksheet.title, grid)
        if sheet:
            sheets.append(sheet)
    workbook.close()
    return sheets


def read_csv(fh) -> list[Sheet]:
    data = fh.read()
    for encoding in CSV_ENCODINGS:
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ReadError("无法识别 CSV 的文字编码（支持 UTF-8 与 GB18030）。")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    grid = [
        [Cell(value, value.strip(), f"{get_column_letter(col_no)}{row_no}")
         for col_no, value in enumerate(row, start=1)]
        for row_no, row in enumerate(csv.reader(io.StringIO(text), dialect), start=1)
    ]
    sheet = _build_sheet("CSV", grid)
    return [sheet] if sheet else []


def _cell(cell, row_no: int, col_no: int) -> Cell:
    coord = f"{get_column_letter(col_no)}{row_no}"
    value = getattr(cell, "value", None)
    number_format = getattr(cell, "number_format", "") or ""
    return Cell(value, _as_text(value), coord, "" if number_format == "General" else number_format)


def _as_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat(" ")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float):
        return format(Decimal(repr(value)).normalize(), "f")  # 80.2 not 80.20000000000000284
    return str(value).strip()


def _build_sheet(name: str, grid: list[list[Cell]]) -> Sheet | None:
    """grid[i] is file row i + 1. Picks the most header-like of the first filled rows."""
    filled = [i for i, row in enumerate(grid) if _filled(row) >= MIN_CELLS_PER_ROW]
    if not filled:
        return None
    header_index = max(filled[:HEADER_SCAN_ROWS], key=lambda i: (_text_cells(grid[i]), -i))
    headers = _unique_headers(grid[header_index])
    rows, skipped = [], []
    for index in range(header_index + 1, len(grid)):
        row, row_no = grid[index], index + 1
        count = _filled(row)
        if count == 0:
            continue
        if count < MIN_CELLS_PER_ROW:
            text = next(c.text for c in row if c.text)
            skipped.append((row_no, f"只有一个单元格有内容（{text[:40]}），视为标题或备注"))
            continue
        rows.append(Row(row_no, {h: row[i] for i, h in enumerate(headers) if i < len(row)}))
    # Columns with no header text are kept only if some row has data in them.
    used = {h for r in rows for h, c in r.cells.items() if c.text}
    header_cells = grid[header_index]
    keep = [h for h, cell in zip(headers, header_cells, strict=False) if cell.text or h in used]
    for r in rows:
        r.cells = {h: c for h, c in r.cells.items() if h in keep}
    return Sheet(name, header_index + 1, keep, rows, skipped)


def _filled(row: list[Cell]) -> int:
    return sum(1 for c in row if c.text)


def _text_cells(row: list[Cell]) -> int:
    """How header-like a row is: filled cells holding text, not numbers or dates."""
    return sum(1 for c in row if c.text and isinstance(c.value, str) and not _numeric(c.text))


def _numeric(text: str) -> bool:
    try:
        float(text.replace(",", ""))
    except ValueError:
        return False
    return True


def _unique_headers(cells: list[Cell]) -> list[str]:
    """Header texts; blanks become "列G", repeats get " (2)", " (3)"."""
    headers, seen = [], {}
    for index, cell in enumerate(cells, start=1):
        text = cell.text or ""
        if not text:
            headers.append(f"列{get_column_letter(index)}")
            continue
        seen[text] = seen.get(text, 0) + 1
        headers.append(text if seen[text] == 1 else f"{text} ({seen[text]})")
    return headers
