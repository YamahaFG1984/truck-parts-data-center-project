"""Quotes for an inquiry (PRD F6, docs/architecture.html §7.6).

unit price = cheapest offer's cost x (1 + margin), Decimal throughout, rounded
half-up to cents (apps.suppliers.services.quoting). Costs and prices are
snapshotted on the QuoteLine so later changes never alter a quote already sent.
"""

import datetime as dt
from decimal import Decimal
from io import BytesIO

from django.conf import settings
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image

from apps.catalog.models import Part
from apps.catalog.services.matcher import alternatives
from apps.suppliers.services.quoting import best_offer, suggested_price

from ..models import Inquiry, QuoteLine

THUMB_PX = 72
CUSTOMER_COLUMNS = [
    ("Item", 6), ("Image", 12), ("SKU", 16), ("Description", 40), ("OE / Cross No.", 24),
    ("Specification", 28), ("Packing", 22), ("Qty", 8), ("Unit Price (USD)", 14),
    ("Amount (USD)", 14), ("MOQ", 8), ("Lead Time", 11), ("Alternatives", 20), ("Note", 24),
]
INTERNAL_COLUMNS = [
    ("SKU", 16), ("Supplier", 30), ("Supplier P/N", 14), ("Unit Cost (USD)", 14),
    ("Margin", 9), ("Unit Price (USD)", 14), ("Quoted On", 12),
]


class QuoteError(ValueError):
    pass


def build_quote(inquiry: Inquiry, part: Part, qty: int, margin: Decimal | None = None) -> QuoteLine:
    if qty <= 0:
        raise QuoteError("数量必须大于 0。")
    margin = settings.DEFAULT_MARGIN if margin is None else Decimal(margin)
    if not Decimal("0") <= margin < Decimal("10"):
        raise QuoteError("毛利率应在 0% 到 1000% 之间。")
    offer = best_offer(part)
    if offer is None:
        raise QuoteError(f"{part.sku} 还没有供应商报价，无法计算售价。")
    note = f"低于起订量（MOQ {offer.moq}）" if offer.moq and qty < offer.moq else ""
    line = QuoteLine.objects.create(
        inquiry=inquiry, part=part, offer=offer, qty=qty, unit_cost_usd=offer.unit_cost_usd,
        margin=margin, unit_price_usd=suggested_price(offer.unit_cost_usd, margin), note=note,
    )
    inquiry.status = Inquiry.Status.QUOTED
    inquiry.save(update_fields=["status", "updated_at"])
    return line


def quote_to_xlsx(inquiry: Inquiry, *, include_costs: bool) -> bytes:
    lines = list(
        inquiry.quote_lines.select_related("part__category", "offer__supplier")
        .prefetch_related("part__numbers", "part__images")
    )
    wb = Workbook()
    ws = wb.active
    ws.title = "Quotation"
    ws.append(["Demo Truck Parts Co., Ltd. — QUOTATION（演示数据，非真实报价）"])
    ws.append([f"Inquiry #{inquiry.pk}    Date: {dt.date.today():%Y-%m-%d}    "
               f"Customer: {inquiry.customer or '-'}    Currency: USD    Term: FOB"])
    ws["A1"].font = Font(bold=True, size=14)
    _header(ws, 4, CUSTOMER_COLUMNS)

    total = Decimal("0")
    for i, line in enumerate(lines, start=1):
        part, offer = line.part, line.offer
        numbers = part.numbers.all()
        row = 4 + i
        ws.append([
            i, "", part.sku, part.title_en or part.name_en,
            "\n".join(f"{n.get_kind_display().split()[-1]} {n.number}"
                      for n in numbers if n.kind in ("OE", "CROSS")),
            "\n".join(f"{k}: {v}" for k, v in part.attributes.items()),
            _packing(part.packaging), line.qty, float(line.unit_price_usd),
            float(line.amount_usd), offer.moq if offer else None,
            f"{offer.lead_days} days" if offer and offer.lead_days else "",
            ", ".join(alternatives(part).values_list("sku", flat=True)[:5]), line.note,
        ])
        total += line.amount_usd
        ws.row_dimensions[row].height = THUMB_PX * 0.78
        for cell in ws[row]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        _add_thumbnail(ws, part, f"B{row}")
    ws.append([])
    ws.append([""] * 8 + ["Total", float(total)])
    ws.cell(ws.max_row, 9).font = ws.cell(ws.max_row, 10).font = Font(bold=True)
    for col in ("I", "J"):
        for cell in ws[col][4:]:
            cell.number_format = "#,##0.00"

    if include_costs:
        internal = wb.create_sheet("Internal")
        _header(internal, 1, INTERNAL_COLUMNS)
        for line in lines:
            offer = line.offer
            internal.append([
                line.part.sku,
                offer.supplier.name if offer else "",
                offer.supplier_pn if offer else "",
                float(line.unit_cost_usd), float(line.margin), float(line.unit_price_usd),
                offer.quoted_at.isoformat() if offer else "",
            ])
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _header(ws, row: int, columns):
    for col, (title, width) in enumerate(columns, start=1):
        cell = ws.cell(row=row, column=col, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D4ED8")
        ws.column_dimensions[cell.column_letter].width = width


def _packing(packaging: dict) -> str:
    parts = []
    if packaging.get("pcs_per_carton"):
        parts.append(f"{packaging['pcs_per_carton']} {packaging.get('unit', 'pc')}/ctn")
    if packaging.get("carton_l_cm"):
        parts.append(f"{packaging['carton_l_cm']}x{packaging.get('carton_w_cm')}x"
                     f"{packaging.get('carton_h_cm')} cm")
    if packaging.get("gross_weight_kg"):
        parts.append(f"G.W. {packaging['gross_weight_kg']} kg")
    return "\n".join(parts)


def _add_thumbnail(ws, part: Part, anchor: str) -> None:
    image = next((i for i in part.images.all() if i.is_primary), None)
    if image is None:
        return
    try:
        with image.image.open("rb") as fh, Image.open(fh) as src:
            thumb = src.convert("RGB")
            thumb.thumbnail((THUMB_PX * 2, THUMB_PX))
            buf = BytesIO()
            thumb.save(buf, "PNG")
    except (OSError, ValueError):
        return  # a missing file must not break the quote
    buf.seek(0)
    ws.add_image(XLImage(buf), anchor)
