"""The three exports of docs/archive-design.html §11, as xlsx workbooks. Every row can be
traced to a row of an archived original; prices keep their own currency and are never
added up or converted."""

import io
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from apps.sources.models import SourceRecord

from ..models import Membership, Product, ReviewItem
from . import products

FILES = {"master": "主数据.xlsx", "quotes": "供应商报价.xlsx", "review": "待人工确认清单.xlsx"}
NOTICE = "演示数据均为合成；价格按供应商原币种列出，未换算、未加总。"


def master() -> Workbook:
    rows = []
    for product in (Product.objects.filter(memberships__isnull=False).distinct()
                    .order_by("id")):
        members = products.members(product)
        lines = {line.field: line for line in products.summary(product, members)}
        rows.append([
            product.code, product.get_status_display(),
            *[_summary_text(lines[name]) for name, _ in products.SUMMARY],
            len(members),
            "；".join(f"{m.supplier.name}: {m.current_record.supplier_sku or '（无料号）'}"
                     for m in members),
            "、".join(line.label for line in lines.values() if line.conflict),
            "、".join(products.missing(product, members)),
            "；".join(m.get_status_display() + " " + m.record_key for m in members
                     if m.status != Membership.Status.ACTIVE),
            "；".join(_where(m.current_record) for m in members),
        ])
    return _book("主数据", ["产品编码", "状态", *[label for _, label in products.SUMMARY],
                          "成员数", "各供应商料号", "冲突字段", "待补充字段", "暂停的成员",
                          "成员出处（供应商 · 记录 · 文件 · 定位）"], rows)


def quotes() -> Workbook:
    product_of = {(m.supplier_id, m.record_key): m.product.code
                  for m in Membership.objects.select_related("product")}
    records = (SourceRecord.objects.current().select_related("supplier", "source_file")
               .order_by("supplier__name", "record_key"))
    rows = [[
        product_of.get((r.supplier_id, r.record_key), ""), r.supplier.name, r.record_key,
        r.supplier_sku, r.name, r.price, r.currency, r.moq, r.quote_date, r.version,
        r.get_change_type_display(), r.source_file.original_name,
        f"第 {r.page} 页 表 {r.table_no}" if r.page else r.sheet, r.row_no, r.locator,
    ] for r in records]
    return _book("供应商报价", ["产品编码", "供应商", "记录 ID", "供应商料号", "原始名称",
                            "单价", "币种", "MOQ", "报价日期", "版本", "变化类型", "原始文件",
                            "工作表 / 页", "行号", "定位"], rows)


def review_list() -> Workbook:
    items = (ReviewItem.objects.filter(status=ReviewItem.Status.OPEN)
             .select_related("record_a__supplier", "record_a__source_file",
                             "record_b__supplier", "record_b__source_file")
             .order_by("category", "id"))
    rows = [[
        item.pk, item.get_kind_display(), item.get_category_display(), item.strength,
        _where(item.record_a), _where(item.record_b) if item.record_b else "",
        "；".join(item.triggers),
        "；".join(f"{c['label']}：{c['a']['value']} ≠ {c['b']['value']}" for c in item.conflicts),
        "；".join(f"{m['label']}（{_side(m['side'])}）" for m in item.missing),
        item.suggested_action, item.get_status_display(),
    ] for item in items]
    return _book("待人工确认", ["条目", "类型", "类别", "强度", "记录 A", "记录 B",
                            "触发原因", "冲突字段（A ≠ B）", "缺失字段", "建议动作", "状态"],
                 rows)


BUILDERS = {"master": master, "quotes": quotes, "review": review_list}


def to_bytes(book: Workbook) -> bytes:
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def write_all(directory) -> list[Path]:
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    paths = []
    for kind, name in FILES.items():
        path = directory / name
        BUILDERS[kind]().save(path)
        paths.append(path)
    return paths


def _book(title, headers, rows) -> Workbook:
    book = Workbook()
    sheet = book.active
    sheet.title = title
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    for i, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = max(10, min(40, len(header) * 3))
    notes = book.create_sheet("说明")
    notes.append([NOTICE])
    return book


def _summary_text(line) -> str:
    if line.conflict:
        return "冲突：" + " / ".join(line.texts)
    return line.text


def _where(record) -> str:
    return (f"{record.supplier.name} · {record.record_key} v{record.version} · "
            f"{record.source_file.original_name} · {record.locator}")


def _side(side) -> str:
    return {"a": "A 缺", "b": "B 缺", "both": "两边都缺"}.get(side, side)
