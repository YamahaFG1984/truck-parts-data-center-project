"""The three exports and the import/export commands (docs/archive-design.html §11)."""

import io
from decimal import Decimal
from io import StringIO

import pytest
from django.core.management import CommandError, call_command
from openpyxl import load_workbook

from apps.archive.models import Membership, Product, ReviewItem
from apps.archive.services import export, review
from apps.sources.models import SourceFile, SourceRecord
from apps.sources.tests.dataset import workbook_x, workbook_x2

from .conftest import find_item

pytestmark = pytest.mark.django_db


def rows(book):
    """Read back what the downloaded file holds."""
    sheet = load_workbook(io.BytesIO(export.to_bytes(book))).worksheets[0]
    header, *body = [list(r) for r in sheet.iter_rows(values_only=True)]
    return [dict(zip(header, r, strict=True)) for r in body]


def test_master_data_has_one_row_per_product(archive, user):
    review.confirm_same(find_item("X-001~Y-001"), user, "同一件")
    data = rows(export.master())

    assert len(data) == Product.objects.filter(memberships__isnull=False).distinct().count()
    grouped = next(r for r in data if r["成员数"] == 2)
    assert grouped["状态"] == "已确认归一" and grouped["品类"] == "Front Grille"
    assert "Supplier X: X-01-00" in grouped["各供应商料号"]
    assert "Supplier X.xlsx · Price List!R2" in grouped["成员出处（供应商 · 记录 · 文件 · 定位）"]


def test_master_data_marks_conflicts_and_missing_fields(archive, user):
    item = find_item("X-008~Y-005")  # left vs right, shared OE
    review.confirm_same(item, user, "演示：人工判为同一件后冲突仍可见")
    data = rows(export.master())
    merged = next(r for r in data if r["成员数"] == 2)
    assert merged["冲突字段"] == "位置" and merged["位置"].startswith("冲突：")
    x009 = Membership.objects.get(record_key="X-009").product.code
    assert "单价" in next(r for r in data if r["产品编码"] == x009)["待补充字段"]


def test_quotes_keep_each_currency_and_point_to_the_original(archive):
    data = rows(export.quotes())

    assert len(data) == SourceRecord.objects.current().count() == 22
    y001 = next(r for r in data if r["记录 ID"] == "Y-001")
    assert (y001["单价"], y001["币种"]) == (46.58, "GBP")
    assert (y001["原始文件"], y001["工作表 / 页"], y001["行号"]) == (
        "Supplier Y.xlsx", "Catalog Export", 2)
    x009 = next(r for r in data if r["记录 ID"] == "X-009")
    assert x009["单价"] is None and x009["币种"] is None  # never defaulted to USD
    assert all(r["产品编码"] for r in data)


def test_the_review_list_has_every_open_item(archive):
    data = rows(export.review_list())
    assert len(data) == ReviewItem.objects.filter(status="open").count()
    conflict = next(r for r in data if r["类别"] == "位置冲突")
    assert "≠" in conflict["冲突字段（A ≠ B）"] and "Price List!R" in conflict["记录 A"]
    assert all(r["建议动作"] for r in data)


def test_archive_export_writes_three_files(archive, tmp_path):
    out = StringIO()
    call_command("archive_export", "--out", str(tmp_path / "out"), stdout=out)
    names = sorted(p.name for p in (tmp_path / "out").iterdir())
    assert names == sorted(export.FILES.values())
    book = load_workbook(tmp_path / "out" / "供应商报价.xlsx")
    assert book.sheetnames == ["供应商报价", "说明"]


@pytest.fixture
def admin(django_user_model):
    return django_user_model.objects.create_superuser("admin", password="x")


def test_archive_import_previews_by_default(admin, tmp_path):
    path = tmp_path / "Supplier X.xlsx"
    path.write_bytes(workbook_x())
    out = StringIO()
    call_command("archive_import", str(path), "--supplier", "Supplier X", stdout=out)

    assert "新增 11" in out.getvalue() and "预检结束，未入库" in out.getvalue()
    assert not SourceRecord.objects.exists()
    assert SourceFile.objects.get().status == "previewed"


def test_archive_import_commits_and_then_shows_changes(admin, tmp_path):
    first, second = tmp_path / "x1.xlsx", tmp_path / "x2.xlsx"
    first.write_bytes(workbook_x())
    second.write_bytes(workbook_x2())
    call_command("archive_import", str(first), "--supplier", "Supplier X", "--commit",
                 stdout=StringIO())
    assert SourceRecord.objects.count() == 11 and Membership.objects.count() == 11

    out = StringIO()
    call_command("archive_import", str(second), "--supplier", "Supplier X", "--commit",
                 stdout=out)
    assert "位置 Left → Right" in out.getvalue() and "已入库 4 条" in out.getvalue()
    assert SourceRecord.objects.current().get(record_key="X-001").price == Decimal("45")


def test_archive_import_accepts_mapping_overrides(admin, tmp_path):
    path = tmp_path / "x.xlsx"
    path.write_bytes(workbook_x())
    out = StringIO()
    call_command("archive_import", str(path), "--supplier", "S", "--map", "Position=ignore",
                 stdout=out)
    assert "Position" in out.getvalue() and "人工指定" in out.getvalue()
    with pytest.raises(CommandError, match="没有这些表头"):
        call_command("archive_import", str(path), "--supplier", "S", "--map", "Nope=name",
                     stdout=StringIO())
